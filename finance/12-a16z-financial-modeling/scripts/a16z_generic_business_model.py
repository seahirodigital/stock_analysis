import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
import os
import sys
from datetime import datetime

def generate_generic_financial_model(output_dir, business_name="Generic_Business"):
    print(f"a16z Financial Modeling: {business_name}向けの汎用Excelモデル(P&L/CF)を構築中...")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3Y_Financial_Model_a16z"

    # --- 1. Assumptions (前提条件) ---
    ws['A1'] = "【前提条件 (Assumptions)】"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = "※ ここ（B列・E列）の数値を変更すると、下のシミュレーションが自動で再計算されます。"
    
    # Revenue Drivers
    ws['A4'] = "[売上・KPI変数]"
    ws['A4'].font = Font(bold=True)
    ws['A5'] = "事業開始月（Month X）"
    ws['B5'] = 1             # 1ヶ月目開始
    ws['A6'] = "初期獲得顧客数/販売数 (月間)"
    ws['B6'] = 100           # 開始月のユーザー数
    ws['A7'] = "月次成長率 (MoM Growth)"
    ws['B7'] = 0.10          # 10%成長
    ws['B7'].number_format = '0.0%'
    ws['A8'] = "平均単価 (ARPU等)"
    ws['B8'] = 10000         # 単価
    ws['A9'] = "売上原価率 (COGS %)"
    ws['B9'] = 0.20          # 20%
    ws['B9'].number_format = '0.0%'

    # Cost Drivers (SGA / CAPEX)
    ws['D4'] = "[コスト・投資変数]"
    ws['D4'].font = Font(bold=True)
    ws['D5'] = "初期投資額 (CAPEX)"
    ws['E5'] = 5000000
    ws['D6'] = "月間マーケティング予算"
    ws['E6'] = 300000
    ws['D7'] = "月間地代家賃・オフィス代"
    ws['E7'] = 200000
    ws['D8'] = "月間システム・サーバー代"
    ws['E8'] = 50000
    
    # Labor Drivers
    ws['G4'] = "[人件費・組織変数]"
    ws['G4'].font = Font(bold=True)
    ws['G5'] = "役員・正社員数"
    ws['H5'] = 2
    ws['G6'] = "平均月給程"
    ws['H6'] = 300000
    ws['G7'] = "外注・アルバイト月間枠"
    ws['H7'] = 100000
    ws['G8'] = "人件費合計 (月)"
    ws['H8'] = "=H5*H6+H7"

    # Number formats
    number_format_cells = ['B6', 'B8', 'E5', 'E6', 'E7', 'E8', 'H6', 'H7', 'H8']
    for cell in number_format_cells:
        ws[cell].number_format = '#,##0'

    # --- 2. P&L / CF Model ---
    start_row = 12
    ws.cell(row=start_row, column=1, value="勘定科目 (P&L Items)").font = Font(bold=True)
    
    months = 36
    for m in range(1, months + 1):
        col = m + 1
        ws.cell(row=start_row, column=col, value=f"Month {m}").font = Font(bold=True)
        ws.cell(row=start_row, column=col).alignment = Alignment(horizontal='right')

    # Rows Definition
    rows_def = [
        ("【主要KPI】", None),
        ("  └ ボリューム (顧客数/販売数等)", "kpi_vol"),
        ("【売上高】", None),
        ("  └ メイン売上", "rev_main"),
        ("売上高 合計", "rev_total"),
        ("【売上原価（変動費）】", None),
        ("  └ COGS (原価率連動)", "cogs"),
        ("売上総利益 (粗利 - Gross Profit)", "gross_profit"),
        ("【販売費及び一般管理費】", None),
        ("  └ 広告宣伝費・マーケティング", "ad"),
        ("  └ 人件費 (組織)", "labor"),
        ("  └ 地代家賃等", "rent"),
        ("  └ ツール・サーバー管理費", "sys"),
        ("販管費 合計", "sga_total"),
        ("【営業利益 (EBITDA)】", "ebitda"),
        ("営業利益率", "ebitda_margin"),
        ("【投資CF・事業CF】", None),
        ("  └ 初期投資 (CAPEX)", "capex"),
        ("当月フリーキャッシュフロー", "fcf"),
        ("累積キャッシュフロー (必要資金)", "cum_cf"),
    ]

    r_idx = start_row + 1
    row_map = {}
    for label, key in rows_def:
        cell = ws.cell(row=r_idx, column=1, value=label)
        if not key:
            cell.font = Font(bold=True)
        if key:
            row_map[key] = r_idx
        r_idx += 1

    ws.column_dimensions['A'].width = 32
    for col in range(2, months + 2):
        ws.column_dimensions[get_column_letter(col)].width = 16

    def cr(r, c): return f"{get_column_letter(c)}{r}"

    for m in range(1, months + 1):
        c = m + 1
        
        # ボリューム推移（単月成長）
        if m == 1:
            ws.cell(row=row_map["kpi_vol"], column=c, value=f"=IF({m}>=$B$5, $B$6, 0)")
        else:
            prev_c = get_column_letter(c - 1)
            ws.cell(row=row_map["kpi_vol"], column=c, value=f"=IF({m}>=$B$5, {prev_c}{row_map['kpi_vol']}*(1+$B$7), 0)")

        # 売上
        ws.cell(row=row_map["rev_main"], column=c, value=f"={cr(row_map['kpi_vol'], c)} * $B$8")
        ws.cell(row=row_map["rev_total"], column=c, value=f"={cr(row_map['rev_main'], c)}")

        # 原価、粗利
        ws.cell(row=row_map["cogs"], column=c, value=f"={cr(row_map['rev_total'], c)} * $B$9")
        ws.cell(row=row_map["gross_profit"], column=c, value=f"={cr(row_map['rev_total'], c)} - {cr(row_map['cogs'], c)}")

        # 販管費 (売上が立つ前でも固定費は開始月から発生と仮定)
        ws.cell(row=row_map["ad"], column=c, value=f"=IF({m}>=$B$5, $E$6, 0)")
        ws.cell(row=row_map["labor"], column=c, value=f"=IF({m}>=$B$5, $H$8, 0)")
        ws.cell(row=row_map["rent"], column=c, value=f"=IF({m}>=$B$5, $E$7, 0)")
        ws.cell(row=row_map["sys"], column=c, value=f"=IF({m}>=$B$5, $E$8, 0)")
        
        ws.cell(row=row_map["sga_total"], column=c, value=f"=SUM({cr(row_map['ad'], c)}:{cr(row_map['sys'], c)})")

        # 営業利益
        ws.cell(row=row_map["ebitda"], column=c, value=f"={cr(row_map['gross_profit'], c)} - {cr(row_map['sga_total'], c)}")
        ws.cell(row=row_map["ebitda_margin"], column=c, value=f"=IF({cr(row_map['rev_total'], c)}>0, {cr(row_map['ebitda'], c)}/{cr(row_map['rev_total'], c)}, 0)")
        
        # FCF
        ws.cell(row=row_map["capex"], column=c, value=f"=IF({m}=$B$5, $E$5, 0)")
        ws.cell(row=row_map["fcf"], column=c, value=f"={cr(row_map['ebitda'], c)} - {cr(row_map['capex'], c)}")

        # 累積キャッシュフロー (ここで一番凹んだ額が「必要調達額」となる)
        if m == 1:
            ws.cell(row=row_map["cum_cf"], column=c, value=f"={cr(row_map['fcf'], c)}")
        else:
            prev_c = get_column_letter(c - 1)
            ws.cell(row=row_map["cum_cf"], column=c, value=f"={prev_c}{row_map['cum_cf']} + {cr(row_map['fcf'], c)}")

    # 表示形式設定
    comma_rows = [v for k, v in row_map.items() if k not in ("kpi_vol", "ebitda_margin")]
    for r in comma_rows:
        for c in range(2, months + 2):
            ws.cell(row=r, column=c).number_format = '#,##0'

    for c in range(2, months + 2):
        ws.cell(row=row_map["kpi_vol"], column=c).number_format = '#,##0'
        ws.cell(row=row_map["ebitda_margin"], column=c).number_format = '0.0%'

    # 太字スタイル
    bold_rows = [row_map["rev_total"], row_map["gross_profit"], row_map["sga_total"], row_map["ebitda"], row_map["fcf"], row_map["cum_cf"]]
    for r in bold_rows:
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in range(2, months + 2):
            ws.cell(row=r, column=c).font = Font(bold=True)

    # --- 3. Save ---
    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M')
    # Use business_name in file name for clarity, defaulting to Generic
    safe_name = "".join([c for c in business_name if c.isalpha() or c.isdigit() or c==' ']).rstrip()
    if not safe_name: safe_name = "Generic"
    
    excel_filename = os.path.join(output_dir, f"{timestamp}_a16z_FinancialModel_{safe_name}.xlsx")
    
    wb.save(excel_filename)
    print(f"汎用財務モデル（Excel・数式連動型）を出力しました。\n変数を書き換えるだけでリアルタイムシミュレーションが可能です。\n出力先: {excel_filename}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python a16z_generic_business_model.py <output_dir> [business_name]")
        print("Outputting to current directory by default as fallback.")
        generate_generic_financial_model(".", "GenericBusiness")
    else:
        out_dir = sys.argv[1]
        b_name = sys.argv[2] if len(sys.argv) > 2 else "GenericBusiness"
        generate_generic_financial_model(out_dir, b_name)
